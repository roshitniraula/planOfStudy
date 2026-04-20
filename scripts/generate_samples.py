"""
generate_samples.py — creates 20 sample Plan of Study .docx files in 'Plan of Study/'.
Run from repo root: python scripts/generate_samples.py
"""

from __future__ import annotations
import random
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_DIR = Path(__file__).parent.parent / "Plan of Study"

# ── Sample data pools ──────────────────────────────────────────────────────────

STUDENTS = [
    ("Emma Johnson",     "Junior",   "Biology",              "Pre-Health Certificate",    "Spring 2027"),
    ("Liam Patel",       "Senior",   "Computer Science",     "Mathematics Minor",          "Spring 2026"),
    ("Sofia Hernandez",  "Sophomore","Psychology",           "Sociology Minor",            "Fall 2027"),
    ("Noah Williams",    "Junior",   "Nursing",              "Spanish Minor",              "Spring 2027"),
    ("Ava Thompson",     "Senior",   "Business Administration","Economics Minor",           "Spring 2026"),
    ("Elijah Davis",     "Junior",   "Political Science",    "History Minor",              "Fall 2027"),
    ("Mia Anderson",     "Sophomore","Biochemistry",         "Chemistry Minor",            "Spring 2028"),
    ("James Wilson",     "Senior",   "Mechanical Engineering","Mathematics Minor",          "Spring 2026"),
    ("Charlotte Moore",  "Junior",   "English",              "Journalism Certificate",     "Spring 2027"),
    ("Benjamin Taylor",  "Sophomore","Environmental Science", "Geography Minor",            "Fall 2027"),
    ("Amelia Jackson",   "Senior",   "Social Work",          "Psychology Minor",           "Spring 2026"),
    ("Lucas Martinez",   "Junior",   "Music",                "Theatre Minor",              "Spring 2027"),
    ("Harper Lee",       "Senior",   "Finance",              "Accounting Minor",           "Fall 2026"),
    ("Henry White",      "Freshman", "Undecided",            None,                         "Spring 2028"),
    ("Evelyn Harris",    "Junior",   "Criminal Justice",     "Sociology Minor",            "Spring 2027"),
    ("Alexander Clark",  "Senior",   "Civil Engineering",    "Environmental Studies Minor","Spring 2026"),
    ("Abigail Lewis",    "Sophomore","Art",                   "Art History Minor",          "Fall 2027"),
    ("Michael Robinson", "Junior",   "Kinesiology",          "Nutrition Certificate",      "Spring 2027"),
    ("Ella Walker",      "Senior",   "Communications",       "Marketing Minor",            "Spring 2026"),
    ("Daniel Hall",      "Junior",   "Philosophy",           "Religious Studies Minor",    "Fall 2027"),
]

YEARS = ["Freshman", "Sophomore", "Junior", "Senior"]

LEADERSHIP_EXPERIENCES = [
    ("Student Government Senator",    "Leadership",   "Completed"),
    ("Resident Advisor",              "Leadership",   "Completed"),
    ("MavPass Mentor",                "Development",  "Completed"),
    ("Honors Student Ambassador",     "Leadership",   "In Progress"),
    ("Club President – STEM Society", "Leadership",   "Completed"),
    ("Peer Tutor Coordinator",        "Development",  "In Progress"),
    ("Orientation Leader",            "Application",  "Completed"),
    ("Team Captain – Soccer",         "Leadership",   "Completed"),
    ("Workshop Facilitator",          "Development",  "Planned"),
    ("Community Garden Organizer",    "Service",      "In Progress"),
]

RESEARCH_EXPERIENCES = [
    ("Undergraduate Research – Biology Lab",    "Research",     "Completed"),
    ("Independent Study: Data Analysis",         "Scholarly",    "In Progress"),
    ("NURS 445 Research Project",               "Research",     "Planned"),
    ("Honors Thesis Proposal",                  "Scholarly",    "In Progress"),
    ("Conference Presentation – NCUR",          "Creative",     "Completed"),
    ("ENG 271W Writing Portfolio",              "Creative",     "Completed"),
    ("Survey Study on Campus Mental Health",    "Research",     "In Progress"),
    ("Archival Research – History Dept",        "Scholarly",    "Planned"),
    ("Engineering Design Capstone",             "Research",     "Completed"),
    ("Creative Writing Collection",             "Creative",     "Completed"),
]

INTERCULTURAL_EXPERIENCES = [
    ("Study Abroad – Spain",                    "Global",       "Completed"),
    ("International Culture Night Volunteer",   "Development",  "Completed"),
    ("Spanish Conversation Partner Program",    "Development",  "In Progress"),
    ("ISSS Global Buddy Program",               "Application",  "Completed"),
    ("Multicultural Affairs Leadership Summit", "Leadership",   "Completed"),
    ("Attend Three Culture Nights",             "Development",  "Completed"),
    ("Study Abroad – Japan",                    "Global",       "Planned"),
    ("Cultural Competency Workshop Series",     "Development",  "In Progress"),
    ("ESL Tutoring Program",                    "Service",      "Completed"),
    ("Global Issues Seminar",                   "Scholarly",    "Completed"),
]

LEADERSHIP_REFLECTIONS = [
    "Your values and how they showed up in your experience",
    "Your leadership strengths and weaknesses and how they showed up",
    "The success and failures you encountered while working in a team",
    "The roles people played and how those roles aligned with their strengths",
    "Different effective and ineffective leadership styles observed",
    "Your leadership philosophy and how it has evolved over time",
]

RESEARCH_REFLECTIONS = [
    "The stages of the research process as they relate to your discipline",
    "The ethical considerations involved in your research",
    "How your findings contribute to existing knowledge",
    "The methodologies you applied and why",
    "How this experience shaped your academic identity",
]

INTERCULTURAL_REFLECTIONS = [
    "Your understanding of your own cultural story, background, and identity",
    "How intercultural communication affected your experience",
    "Challenges and growth from engaging across cultural differences",
    "How this experience broadened your global perspective",
    "The relationship between culture and your professional field",
]

ADVISOR_NOTES = [
    "Great progress this semester. Continue developing your reflection on the leadership experience before HONR 375.",
    "Needs a formal proposal submitted before this experience can be approved. Schedule advising appointment.",
    "Strong engagement with the competency. Consider adding one more highlighted experience to strengthen your portfolio.",
    "In HONR 201, write your practice reflection on this experience. Focus on growth and challenges.",
    "This student is on track for graduation. All three competency areas show solid development.",
    None,
    "Please revise the draft reflection — it needs a stronger connection to the required topics listed.",
    "Excellent work on the research project. The thesis proposal is well-scoped for the timeline.",
    None,
    "Reminder: intercultural experience needs to be completed before Spring registration for HONR 475.",
]

SEMESTERS_PLANNED  = ["Fall 2025", "Spring 2026", "Fall 2026", "Spring 2027", None]
SEMESTERS_COMPLETED = ["Fall 2024", "Spring 2025", "Fall 2025", "Spring 2026", None]


# ── Document builder ───────────────────────────────────────────────────────────

def add_merged_row(table, text: str, bold: bool = False, shade: str | None = None):
    """Add a single-cell merged row to a table."""
    row = table.add_row()
    cell = row.cells[0]
    # Merge across all columns by extending to the last cell
    if len(row.cells) > 1:
        cell = row.cells[0].merge(row.cells[-1])
    cell.text = text
    if bold:
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
    if shade:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), shade)
        tcPr.append(shd)
    return cell


def set_cell(cell, text: str | None):
    cell.text = text or ""


def add_checkbox_paragraphs(cell, stages: list[bool], labels: list[str]):
    """Write checkbox paragraphs into a cell (replaces existing text)."""
    for para in cell.paragraphs:
        para.clear()
    # Remove extra paragraphs
    tc = cell._tc
    for p in tc.findall(qn("w:p"))[1:]:
        tc.remove(p)

    first = True
    for checked, label in zip(stages, labels):
        char = "☒" if checked else "☐"
        if first:
            cell.paragraphs[0].text = f"{char}{label}"
            first = False
        else:
            para = cell.add_paragraph(f"{char}{label}")


def add_reflection_paragraphs(cell, topics: list[str], checked: list[bool]):
    """Write reflection topic checkboxes into a cell."""
    tc = cell._tc
    for p in tc.findall(qn("w:p"))[1:]:
        tc.remove(p)
    first = True
    for i, topic in enumerate(topics):
        char = "☒" if (i < len(checked) and checked[i]) else "☐"
        if first:
            cell.paragraphs[0].text = f"{char} {topic}"
            first = False
        else:
            cell.add_paragraph(f"{char} {topic}")


def build_student_info_table(doc, name, year, date, majors, minors, graduation, portfolio):
    table = doc.add_table(rows=0, cols=6)
    # Row 0: header
    add_merged_row(table, "Student Information", bold=True, shade="782F40")
    # Row 1: labels
    r1 = table.add_row()
    r1.cells[0].merge(r1.cells[1]).text = "Student Name"
    r1.cells[2].merge(r1.cells[3]).text = "Year"
    r1.cells[4].merge(r1.cells[5]).text = "Date"
    # Row 2: values
    r2 = table.add_row()
    r2.cells[0].merge(r2.cells[1]).text = name or ""
    r2.cells[2].merge(r2.cells[3]).text = year or ""
    r2.cells[4].merge(r2.cells[5]).text = date or ""
    # Row 3: labels
    r3 = table.add_row()
    r3.cells[0].text = "Major(s)"
    r3.cells[1].merge(r3.cells[2]).text = "Minor(s)/Certificate(s)"
    r3.cells[3].merge(r3.cells[5]).text = "Expected Graduation"
    # Row 4: values
    r4 = table.add_row()
    r4.cells[0].text = majors or ""
    r4.cells[1].merge(r4.cells[2]).text = minors or ""
    r4.cells[3].merge(r4.cells[5]).text = graduation or ""
    # Row 5: portfolio label
    add_merged_row(table, "Portfolio Link")
    # Row 6: portfolio value
    add_merged_row(table, portfolio or "")
    return table


def build_required_curriculum_table(doc, honr201_planned, honr201_completed,
                                     honr375_planned, honr375_completed,
                                     honr475_planned, honr475_completed):
    table = doc.add_table(rows=0, cols=4)
    add_merged_row(table, "Required Curriculum", bold=True, shade="782F40")
    rh = table.add_row()
    for i, h in enumerate(["Courses", "Credits", "Semester Planned", "Semester Completed"]):
        rh.cells[i].text = h

    for course, credits, planned, completed in [
        ("HONR 201: Introduction to Honors\nPre-requisite: Admission into Honors\nIn-Person",
         "1", honr201_planned, honr201_completed),
        ("HONR 375: Honors Portfolio Development\nPre-requisite: Complete 4 experiences\nIn-Person",
         "1", honr375_planned, honr375_completed),
        ("HONR 475: Honors Senior Portfolio\nPre-requisite: Complete all 8 experiences\nOnline",
         "1", honr475_planned, honr475_completed),
    ]:
        r = table.add_row()
        r.cells[0].text = course
        r.cells[1].text = credits
        r.cells[2].text = planned or ""
        r.cells[3].text = completed or ""
    return table


def build_optional_curriculum_table(doc):
    table = doc.add_table(rows=0, cols=4)
    add_merged_row(table, "Optional Curriculum * Honors with Distinction*", bold=True)
    add_merged_row(table, "To achieve this distinction, students must complete three language courses beyond the intermediate level.")
    rh = table.add_row()
    for i, h in enumerate(["Courses", "Credits", "Semester Planned", "Semester Completed"]):
        rh.cells[i].text = h
    for _ in range(3):
        table.add_row()
    return table


def build_competency_table(doc, title: str, is_leadership: bool,
                            log_entries: list[tuple],
                            highlighted_entries: list[tuple],
                            advisor_notes: str | None,
                            log_header: str,
                            highlight_header: str,
                            log_col_label: str,
                            reflections: list[str]):
    table = doc.add_table(rows=0, cols=4)

    if is_leadership:
        add_merged_row(table, "HONORS COMPETENCY EXPERIENCES", bold=True, shade="782F40")

    add_merged_row(table, title, bold=True, shade="C8A84B")
    add_merged_row(table, log_header)

    # Experience log header
    rlh = table.add_row()
    rlh.cells[0].merge(rlh.cells[1]).text = log_col_label
    rlh.cells[2].text = "Experience Type"
    rlh.cells[3].text = "Progress"

    # Experience log data (5 rows)
    for i in range(5):
        r = table.add_row()
        if i < len(log_entries):
            exp, typ, prog = log_entries[i]
            r.cells[0].merge(r.cells[1]).text = exp or ""
            r.cells[2].text = typ or ""
            r.cells[3].text = prog or ""
        else:
            r.cells[0].merge(r.cells[1]).text = ""
            r.cells[2].text = ""
            r.cells[3].text = ""

    add_merged_row(table, highlight_header)

    # Highlighted header
    rhh = table.add_row()
    rhh.cells[0].text = "Experience"
    rhh.cells[1].text = "Experience Type"
    rhh.cells[2].text = "Progress Tracker"
    rhh.cells[3].text = "Required Reflection Topics" if is_leadership else "Reflection Topics"

    # Highlighted data (4 rows)
    stage_labels = [
        "Experience is approved through list or proposal",
        "Draft of reflection has been completed",
        "Final draft has been evaluated in HONR 375/475",
        "All remaining revisions have been made",
    ]
    for i in range(4):
        r = table.add_row()
        if i < len(highlighted_entries):
            exp, typ, stages, ref_checked = highlighted_entries[i]
            r.cells[0].text = exp or ""
            r.cells[1].text = typ or ""
            add_checkbox_paragraphs(r.cells[2], stages, stage_labels)
            add_reflection_paragraphs(r.cells[3], reflections, ref_checked)
        else:
            r.cells[0].text = ""
            r.cells[1].text = ""
            add_checkbox_paragraphs(r.cells[2], [False, False, False, False], stage_labels)
            r.cells[3].text = ""

    add_merged_row(table, "Instructor/Advisor Notes", bold=True)
    add_merged_row(table, advisor_notes or "")
    return table


# ── Student scenario generator ─────────────────────────────────────────────────

def random_stages(progress: str, rng: random.Random) -> list[bool]:
    if progress == "Completed":
        n = rng.choice([2, 3, 4])
    elif progress == "In Progress":
        n = rng.choice([1, 2])
    else:
        n = rng.choice([0, 1])
    return [i < n for i in range(4)]


def make_student_doc(index: int) -> None:
    name, year, major, minor, graduation = STUDENTS[index]
    rng = random.Random(index * 42)
    date = rng.choice(["January 2026", "February 2026", "March 2026", "April 2026", None])

    # Curriculum semesters based on year
    if year in ("Freshman", "Sophomore"):
        h201c = rng.choice(["Fall 2024", "Fall 2025"])
        h201p = None
        h375c = None
        h375p = rng.choice(["Spring 2026", "Fall 2026"])
        h475c = None
        h475p = rng.choice(["Spring 2027", "Fall 2027"])
    elif year == "Junior":
        h201c = rng.choice(["Fall 2023", "Fall 2024"])
        h201p = None
        h375c = rng.choice([None, "Spring 2025"])
        h375p = None if h375c else "Spring 2026"
        h475c = None
        h475p = rng.choice(["Spring 2027", "Fall 2026"])
    else:  # Senior
        h201c = "Fall 2023"
        h201p = None
        h375c = rng.choice(["Spring 2024", "Fall 2024"])
        h375p = None
        h475c = None
        h475p = "Spring 2026"

    # Pick experiences per competency
    n_lead = rng.randint(2, 4)
    n_res  = rng.randint(1, 4)
    n_int  = rng.randint(1, 4)

    lead_pool = rng.sample(LEADERSHIP_EXPERIENCES, n_lead)
    res_pool  = rng.sample(RESEARCH_EXPERIENCES,  n_res)
    int_pool  = rng.sample(INTERCULTURAL_EXPERIENCES, n_int)

    def make_highlighted(pool, reflections, max_hl=2):
        hl = []
        for exp, typ, prog in pool[:max_hl]:
            stages = random_stages(prog, rng)
            n_ref = rng.randint(0, len(reflections))
            ref_checked = [rng.random() < 0.5 for _ in range(n_ref)]
            hl.append((exp, typ, stages, ref_checked))
        return hl

    lead_hl = make_highlighted(lead_pool, LEADERSHIP_REFLECTIONS)
    res_hl  = make_highlighted(res_pool,  RESEARCH_REFLECTIONS)
    int_hl  = make_highlighted(int_pool,  INTERCULTURAL_REFLECTIONS)

    doc = Document()

    build_student_info_table(doc, name, year, date, major, minor, graduation, None)
    doc.add_paragraph()

    build_required_curriculum_table(doc, h201p, h201c, h375p, h375c, h475p, h475c)
    doc.add_paragraph()

    build_optional_curriculum_table(doc)
    doc.add_paragraph()

    build_competency_table(
        doc, "LEADERSHIP", True,
        lead_pool, lead_hl,
        rng.choice(ADVISOR_NOTES),
        "Leadership Experience Log: What experiences have you had that have helped you develop as a leader?",
        "In the space below, identify leadership experiences from above to highlight in your Honors Portfolio.",
        "Leadership Experiences and/or Ideas",
        LEADERSHIP_REFLECTIONS,
    )
    doc.add_paragraph()

    build_competency_table(
        doc, "RESEARCH, SCHOLARLY, & CREATIVE ACTIVITY", False,
        res_pool, res_hl,
        rng.choice(ADVISOR_NOTES),
        "Research Experience Log: What experiences have you had that have helped you develop as a researcher or creative?",
        "In the space below, identify research experiences from above to highlight in your Honors Portfolio.",
        "Research Experiences and/or Ideas",
        RESEARCH_REFLECTIONS,
    )
    doc.add_paragraph()

    build_competency_table(
        doc, "INTERCULTURAL ENGAGEMENT", False,
        int_pool, int_hl,
        rng.choice(ADVISOR_NOTES),
        "Intercultural Engagement Experience Log: What experiences have you had that have helped you engage across cultures?",
        "In the space below, identify intercultural engagement experiences from above to highlight in your Honors Portfolio.",
        "Intercultural Experiences and/or Ideas",
        INTERCULTURAL_REFLECTIONS,
    )

    slug = name.lower().replace(" ", "-")
    out = OUT_DIR / f"{slug}.docx"
    doc.save(str(out))
    print(f"  wrote {out.name}")


def main():
    OUT_DIR.mkdir(exist_ok=True)
    print(f"Generating {len(STUDENTS)} sample documents into '{OUT_DIR}'...")
    for i in range(len(STUDENTS)):
        make_student_doc(i)
    print("Done.")


if __name__ == "__main__":
    main()
