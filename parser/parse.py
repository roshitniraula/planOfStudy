"""
parse.py — reads every .docx in 'Plan of Study/' and writes one JSON per student to 'data/'.
"""

import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

PLANS_DIR = Path(__file__).parent.parent / "Plan of Study"
DATA_DIR = Path(__file__).parent.parent / "data"

PLACEHOLDERS = {"Type here.", "Choose an item.", "Click or tap here to enter text."}

PROGRESS_STAGE_LABELS = [
    "Experience is approved through list or proposal",
    "Draft of reflection has been completed",
    "Final draft has been evaluated in HONR 375/475",
    "All remaining revisions have been made",
]


def slugify(name: str) -> str:
    name = name.lower().strip()
    name = re.sub(r"[^\w\s-]", "", name)
    name = re.sub(r"[\s_]+", "-", name)
    name = re.sub(r"-+", "-", name)
    return name.strip("-")


def clean(text: str | None) -> str | None:
    if text is None:
        return None
    text = text.strip()
    if not text or text in PLACEHOLDERS:
        return None
    return text


def cell_text(cell) -> str | None:
    return clean(" ".join(p.text for p in cell.paragraphs).strip()) if cell else None


def raw_cell_text(cell) -> str:
    return "\n".join(p.text for p in cell.paragraphs).strip() if cell else ""


def identify_table(table) -> str | None:
    """Return a table type string based on its header row content."""
    if not table.rows:
        return None
    first_row_text = " ".join(
        c.text.strip().lower() for c in table.rows[0].cells
    )
    if "student name" in first_row_text or "student  name" in first_row_text:
        return "student_info"
    if "honr 201" in first_row_text or (
        "courses" in first_row_text and "credits" in first_row_text and "semester planned" in first_row_text
        and any("honr" in c.text.lower() for row in table.rows for c in row.cells)
    ):
        return "required_curriculum"
    if "honors with distinction" in first_row_text:
        return "optional_curriculum"

    # Check for curriculum tables by column headers
    header_cells = [c.text.strip().lower() for c in table.rows[0].cells]
    if set(header_cells) >= {"courses", "credits"} or (
        "courses" in header_cells and "credits" in header_cells
    ):
        # Could be required or optional — need more context; check subsequent rows
        for row in table.rows[1:]:
            for cell in row.cells:
                t = cell.text.strip().lower()
                if "honr 201" in t or "honr 375" in t or "honr 475" in t:
                    return "required_curriculum"
        return "optional_curriculum"

    # Experience log: has "experience/idea" or "experience type" in header
    if "experience/idea" in first_row_text or "experience type" in first_row_text:
        return "experience_log"

    # Highlighted experiences: has "progress tracker" or "reflection topics" in header
    if "progress tracker" in first_row_text or "reflection topics" in first_row_text:
        return "highlighted"

    return None


def parse_student_info(table) -> dict:
    info = {
        "name": None, "year": None, "date": None,
        "majors": None, "minors_certificates": None,
        "expected_graduation": None, "portfolio_link": None,
    }
    for row in table.rows:
        cells = row.cells
        for i in range(0, len(cells) - 1, 2):
            label = cells[i].text.strip().lower().rstrip(":")
            value = clean(cells[i + 1].text.strip())
            if "student name" in label or label == "name":
                info["name"] = value
            elif "year" in label:
                info["year"] = value
            elif label == "date":
                info["date"] = value
            elif "major" in label:
                info["majors"] = value
            elif "minor" in label or "certificate" in label:
                info["minors_certificates"] = value
            elif "graduation" in label:
                info["expected_graduation"] = value
            elif "portfolio" in label:
                info["portfolio_link"] = value
    return info


def parse_curriculum_table(table) -> list:
    rows = []
    for row in table.rows[1:]:  # skip header
        cells = row.cells
        if len(cells) < 4:
            continue
        course = cell_text(cells[0])
        credits = cell_text(cells[1])
        planned = cell_text(cells[2])
        completed = cell_text(cells[3])
        if course or credits or planned or completed:
            rows.append({
                "course": course,
                "credits": credits,
                "semester_planned": planned,
                "semester_completed": completed,
            })
    return rows


def parse_experience_log(table) -> list:
    entries = []
    for row in table.rows[1:]:  # skip header
        cells = row.cells
        if len(cells) < 3:
            continue
        exp = cell_text(cells[0])
        typ = cell_text(cells[1])
        progress = cell_text(cells[2])
        if exp or typ or progress:
            entries.append({"experience": exp, "type": typ, "progress": progress})
    return entries


def parse_checkboxes(text: str) -> list[bool]:
    results = []
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        if "☒" in line:
            results.append(True)
        elif "☐" in line:
            results.append(False)
    return results


def parse_highlighted(table) -> list:
    entries = []
    for row in table.rows[1:]:  # skip header
        cells = row.cells
        if len(cells) < 4:
            continue
        exp = cell_text(cells[0])
        typ = cell_text(cells[1])

        # Progress tracker (col 2)
        progress_text = raw_cell_text(cells[2])
        progress_stages = parse_checkboxes(progress_text)
        # Pad/truncate to exactly 4
        while len(progress_stages) < 4:
            progress_stages.append(False)
        progress_stages = progress_stages[:4]

        # Reflection topics (col 3)
        reflection_text = raw_cell_text(cells[3])
        reflection_checked = [b for b in parse_checkboxes(reflection_text) if b]

        stages_completed = sum(1 for i, v in enumerate(progress_stages) if v and all(progress_stages[:i + 1]))

        if exp or typ or any(progress_stages):
            entries.append({
                "experience": exp,
                "type": typ,
                "progress_stages_completed": stages_completed,
                "progress_stages": progress_stages,
                "reflection_topics_checked": reflection_checked,
            })
    return entries


def find_advisor_notes(doc, after_table_idx: int) -> str | None:
    """Find the advisor notes paragraph after a given table index."""
    # Advisor notes are free-text paragraphs between highlighted table and next section
    # Iterate document body XML to find paragraphs after the table
    tables_seen = 0
    collect = False
    notes_parts = []
    for elem in doc.element.body:
        tag = elem.tag.split("}")[-1]
        if tag == "tbl":
            tables_seen += 1
            if tables_seen == after_table_idx + 1:
                collect = True
                continue
            if collect:
                break  # hit next table, stop
        if collect and tag == "p":
            text = "".join(r.text for r in elem.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
            if text.strip():
                notes_parts.append(text.strip())
    return clean("\n".join(notes_parts)) if notes_parts else None


def parse_docx(path: Path) -> tuple[dict, list[str]]:
    warnings: list[str] = []
    doc = Document(str(path))

    # Categorize all tables
    tables_by_type: dict[str, list[tuple[int, object]]] = {}
    for idx, table in enumerate(doc.tables):
        ttype = identify_table(table)
        if ttype:
            tables_by_type.setdefault(ttype, []).append((idx, table))
        # else: unrecognized table — could be section heading table, ignore

    # Student info
    student = {
        "name": None, "year": None, "date": None,
        "majors": None, "minors_certificates": None,
        "expected_graduation": None, "portfolio_link": None,
    }
    if "student_info" in tables_by_type:
        student = parse_student_info(tables_by_type["student_info"][0][1])
    else:
        warnings.append("Could not find student information table")

    # Required curriculum
    required = []
    if "required_curriculum" in tables_by_type:
        required = parse_curriculum_table(tables_by_type["required_curriculum"][0][1])
    else:
        warnings.append("Could not find required curriculum table")

    # Optional curriculum
    optional = []
    if "optional_curriculum" in tables_by_type:
        optional = parse_curriculum_table(tables_by_type["optional_curriculum"][0][1])

    # Competency sections — we expect pairs: (experience_log, highlighted) × 3
    # in document order: Leadership, Research, Intercultural
    exp_logs = tables_by_type.get("experience_log", [])
    highlighted_tables = tables_by_type.get("highlighted", [])

    competency_keys = ["leadership", "research", "intercultural"]
    competencies = {}

    for i, key in enumerate(competency_keys):
        log = []
        highlighted = []
        notes = None

        if i < len(exp_logs):
            log = parse_experience_log(exp_logs[i][1])
        else:
            warnings.append(f"Could not find experience log table for {key}")

        if i < len(highlighted_tables):
            highlighted = parse_highlighted(highlighted_tables[i][1])
            # Advisor notes come after highlighted table in document body
            hl_idx = highlighted_tables[i][0]
            notes = find_advisor_notes(doc, hl_idx)
        else:
            warnings.append(f"Could not find highlighted experiences table for {key}")

        competencies[key] = {
            "experience_log": log,
            "highlighted": highlighted,
            "advisor_notes": notes,
        }

    return {
        "source_file": path.name,
        "parsed_at": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "parse_warnings": warnings,
        "student": student,
        "required_curriculum": required,
        "optional_curriculum": optional,
        "competencies": competencies,
    }, warnings


def main():
    DATA_DIR.mkdir(exist_ok=True)
    docx_files = list(PLANS_DIR.glob("*.docx"))

    if not docx_files:
        print("No .docx files found in 'Plan of Study/'. data/ is empty.")
        return

    total = len(docx_files)
    warned = 0

    for path in docx_files:
        try:
            result, warnings = parse_docx(path)
        except Exception as e:
            print(f"ERROR parsing {path.name}: {e}", file=sys.stderr)
            warned += 1
            continue

        name = result["student"].get("name")
        slug = slugify(name) if name else slugify(path.stem)
        out_path = DATA_DIR / f"{slug}.json"

        # Avoid slug collision
        counter = 1
        base_slug = slug
        while out_path.exists():
            slug = f"{base_slug}-{counter}"
            out_path = DATA_DIR / f"{slug}.json"
            counter += 1

        out_path.write_text(json.dumps(result, indent=2), encoding="utf-8")

        if warnings:
            warned += 1
            print(f"  WARN {path.name}: {'; '.join(warnings)}")
        else:
            print(f"  OK   {path.name} -> {out_path.name}")

    clean_count = total - warned
    if warned:
        print(f"\nParsed {clean_count}/{total} files cleanly. {warned} files had warnings.")
    else:
        print(f"\nParsed {total}/{total} files cleanly.")


if __name__ == "__main__":
    main()
